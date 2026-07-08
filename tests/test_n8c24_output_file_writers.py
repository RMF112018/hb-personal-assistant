"""N8C-24 — real generated-file rendering (no faked formats) + atomic write."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from hb_assistant.nas_mcp import client_output_writers as w
from tests.n8c24_helpers import good_zip_b64, make_env


def _cfg(tmp_path: Path):
    return make_env(tmp_path)["config"]


@pytest.mark.parametrize("ft,mode,content", [
    ("txt", "text", "hello"),
    ("md", "markdown_text", "# H\nbody"),
    ("html", "html_text", "<b>x</b>"),
    ("json", "json_text", '{"a":1}'),
    ("csv", "csv_text", "a,b\n1,2"),
])
def test_text_like_render(tmp_path: Path, ft, mode, content) -> None:
    data, extra = w.render_output_bytes(config=_cfg(tmp_path), file_type=ft, content_mode=mode, content=content)
    assert data and isinstance(data, bytes)


def test_docx_is_real_and_readable(tmp_path: Path) -> None:
    from docx import Document
    data, _ = w.render_output_bytes(config=_cfg(tmp_path), file_type="docx",
                                    content_mode="docx_from_markdown_or_text", content="# Title\npara")
    doc = Document(io.BytesIO(data))
    assert any(p.text for p in doc.paragraphs)


def test_xlsx_is_real_and_readable(tmp_path: Path) -> None:
    from openpyxl import load_workbook
    data, _ = w.render_output_bytes(config=_cfg(tmp_path), file_type="xlsx",
                                    content_mode="xlsx_from_csv", content="a,b\n1,2")
    wb = load_workbook(io.BytesIO(data))
    assert wb.active["A1"].value == "a"


def test_pptx_is_real_and_readable(tmp_path: Path) -> None:
    from pptx import Presentation
    data, _ = w.render_output_bytes(config=_cfg(tmp_path), file_type="pptx",
                                    content_mode="pptx_from_markdown_or_json", content="Title\nbody")
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) >= 1


def test_pdf_is_real(tmp_path: Path) -> None:
    data, _ = w.render_output_bytes(config=_cfg(tmp_path), file_type="pdf",
                                    content_mode="pdf_from_html_or_markdown", content="# H\nline")
    assert data[:5] == b"%PDF-"


def test_json_invalid_rejected(tmp_path: Path) -> None:
    with pytest.raises(w.OutputWriteError):
        w.render_output_bytes(config=_cfg(tmp_path), file_type="json", content_mode="json_text", content="{bad")


def test_zip_render_validates(tmp_path: Path) -> None:
    data, extra = w.render_output_bytes(config=_cfg(tmp_path), file_type="zip", content_mode="zip_base64",
                                        content=good_zip_b64())
    assert extra["zip_validation"]["member_count"] == 2


def test_atomic_write_roundtrip(tmp_path: Path) -> None:
    target = tmp_path / "sub" / "x.txt"
    meta = w.write_output_bytes(target, b"payload")
    assert target.read_bytes() == b"payload"
    assert meta["bytes_written"] == 7 and len(meta["sha256"]) == 64


def test_oversized_rejected(tmp_path: Path) -> None:
    cfg = _cfg(tmp_path)
    object.__setattr__(cfg, "max_client_output_file_bytes", 4)
    with pytest.raises(w.OutputWriteError, match="exceeds"):
        w.render_output_bytes(config=cfg, file_type="txt", content_mode="text", content="way too long")
